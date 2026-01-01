# views/components/export_excel.py
"""
Module pour l'export de bulletins en Excel
"""

import pandas as pd
from datetime import datetime
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# Créer le dossier d'exports s'il n'existe pas
EXPORT_DIR = Path("static/exports/excel")
EXPORT_DIR.mkdir(parents=True, exist_ok=True)


def export_bulletin_excel(student, grades_data, semester="Semestre 1", year="2024-2025"):
    """
    Exporte le bulletin d'un étudiant en Excel
    
    Args:
        student: Objet avec attributs username, name, class_name
        grades_data: Liste de dictionnaires {subject, grade, coefficient, teacher}
        semester: Période
        year: Année scolaire
    
    Returns:
        str: Chemin du fichier Excel généré
    """
    
    filename = f"bulletin_{student.username}_{semester.replace(' ', '_')}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    filepath = EXPORT_DIR / filename
    
    # Préparer les données du bulletin
    bulletin_data = []
    total_points = 0
    total_coef = 0
    
    for grade_info in grades_data:
        subject = grade_info.get('subject', 'N/A')
        grade = grade_info.get('grade')
        coef = grade_info.get('coefficient', 1)
        teacher = grade_info.get('teacher', 'N/A')
        
        if grade is not None:
            weighted = grade * coef
            total_points += weighted
            total_coef += coef
            
            bulletin_data.append({
                'Matière': subject,
                'Note': round(grade, 2),
                'Coefficient': coef,
                'Note × Coef': round(weighted, 2),
                'Enseignant': teacher
            })
    
    # Calculer la moyenne
    average = total_points / total_coef if total_coef > 0 else 0
    
    # Créer le DataFrame
    df = pd.DataFrame(bulletin_data)
    
    # Exporter vers Excel avec mise en forme
    with pd.ExcelWriter(filepath, engine='openpyxl') as writer:
        df.to_excel(writer, sheet_name='Bulletin', index=False, startrow=6)
        
        workbook = writer.book
        worksheet = writer.sheets['Bulletin']
        
        # === TITRE ===
        worksheet['A1'] = 'BULLETIN DE NOTES'
        worksheet['A1'].font = Font(size=18, bold=True, color='1E3A8A')
        worksheet['A1'].alignment = Alignment(horizontal='center')
        worksheet.merge_cells('A1:E1')
        
        # === INFORMATIONS ÉTUDIANT ===
        worksheet['A3'] = f'Nom: {student.name}'
        worksheet['A4'] = f'Matricule: {student.username}'
        worksheet['A5'] = f'Classe: {student.class_name}'
        
        worksheet['C3'] = f'Période: {semester}'
        worksheet['C4'] = f'Année: {year}'
        
        for row in [3, 4, 5]:
            worksheet[f'A{row}'].font = Font(size=11)
            worksheet[f'C{row}'].font = Font(size=11)
        
        # === STYLE DE L'EN-TÊTE DES NOTES (ligne 7) ===
        header_fill = PatternFill(start_color='1E3A8A', end_color='1E3A8A', fill_type='solid')
        header_font = Font(color='FFFFFF', bold=True, size=11)
        
        for cell in worksheet[7]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')
        
        # === BORDURES ===
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        for row in worksheet.iter_rows(min_row=7, max_row=len(bulletin_data)+7):
            for cell in row:
                cell.border = thin_border
                cell.alignment = Alignment(horizontal='center', vertical='center')
        
        # Alignement à gauche pour les matières et enseignants
        for row in range(8, len(bulletin_data)+8):
            worksheet[f'A{row}'].alignment = Alignment(horizontal='left', vertical='center')
            worksheet[f'E{row}'].alignment = Alignment(horizontal='left', vertical='center')
        
        # === LIGNE DE MOYENNE ===
        last_row = len(bulletin_data) + 8
        worksheet[f'A{last_row}'] = 'MOYENNE GÉNÉRALE'
        worksheet[f'B{last_row}'] = round(average, 2)
        worksheet[f'C{last_row}'] = total_coef
        worksheet[f'D{last_row}'] = round(total_points, 2)
        worksheet[f'E{last_row}'] = ''
        
        moyenne_fill = PatternFill(start_color='FEF3C7', end_color='FEF3C7', fill_type='solid')
        moyenne_font = Font(bold=True, size=12, color='1E3A8A')
        
        for col in ['A', 'B', 'C', 'D', 'E']:
            cell = worksheet[f'{col}{last_row}']
            cell.fill = moyenne_fill
            cell.font = moyenne_font
            cell.border = thin_border
            cell.alignment = Alignment(horizontal='center', vertical='center')
        
        worksheet[f'A{last_row}'].alignment = Alignment(horizontal='left', vertical='center')
        
        # === MENTION ===
        mention_row = last_row + 2
        mention, _ = get_mention_text(average)
        
        worksheet[f'A{mention_row}'] = f'Mention: {mention}'
        worksheet[f'A{mention_row}'].font = Font(size=14, bold=True, color='1E3A8A')
        worksheet[f'A{mention_row}'].alignment = Alignment(horizontal='center')
        worksheet.merge_cells(f'A{mention_row}:E{mention_row}')
        
        # === AJUSTER LES LARGEURS ===
        worksheet.column_dimensions['A'].width = 30
        worksheet.column_dimensions['B'].width = 12
        worksheet.column_dimensions['C'].width = 12
        worksheet.column_dimensions['D'].width = 15
        worksheet.column_dimensions['E'].width = 25
        
        # === PIED DE PAGE ===
        footer_row = mention_row + 3
        worksheet[f'A{footer_row}'] = f'Document généré le {datetime.now().strftime("%d/%m/%Y à %H:%M")}'
        worksheet[f'A{footer_row}'].font = Font(size=9, italic=True, color='64748B')
        worksheet[f'A{footer_row}'].alignment = Alignment(horizontal='center')
        worksheet.merge_cells(f'A{footer_row}:E{footer_row}')
    
    print(f"✅ Bulletin Excel généré: {filepath}")
    return str(filepath)


def get_mention_text(average):
    """Retourne la mention selon la moyenne"""
    if average >= 16:
        return "Très Bien", "green"
    elif average >= 14:
        return "Bien", "blue"
    elif average >= 12:
        return "Assez Bien", "orange"
    elif average >= 10:
        return "Passable", "grey"
    else:
        return "Insuffisant", "red"